"""Every loader emits the same block model (ARCHITECTURE §5.1). Fixtures are built
here, so tests never depend on the reviewers' ./data."""

import docx
import pytest

from app.ingestion.blocks import Block
from app.ingestion.loaders import load_document
from tests.fixtures.make_pdf import sample_policy_pdf


def _summary(blocks: list[Block]) -> list[tuple]:
    return [(b.kind, b.level, b.text) for b in blocks]


EXPECTED = [
    ("heading", 1, "Test Policy"),
    ("heading", 2, "1. Leave rules"),
    ("text", 0, "Staff get 7 test days each year. This line continues the same paragraph."),
    ("text", 0, "A second paragraph starts here."),
    ("heading", 2, "2. Rates"),
    ("table", 0, "| Band | Days |\n|---|---|\n| A | 1.5 |\n| B | 2.0 |"),
]


def test_pdf_headings_paragraphs_tables_and_no_footer(tmp_path):
    path = tmp_path / "p.pdf"
    path.write_bytes(sample_policy_pdf())
    doc = load_document(str(path))
    assert _summary(doc.blocks) == EXPECTED
    assert {b.page for b in doc.blocks} == {1}
    assert doc.blocks[-1].rows == (("Band", "Days"), ("A", "1.5"), ("B", "2.0"))
    assert not doc.degraded


def test_pdf_blank_page_without_ocr_is_degraded(tmp_path):
    path = tmp_path / "p.pdf"
    path.write_bytes(sample_policy_pdf(with_scanned_page=True))
    doc = load_document(str(path))
    assert _summary(doc.blocks) == EXPECTED
    assert doc.degraded
    assert any("page 2" in w for w in doc.warnings)


def test_docx(tmp_path):
    d = docx.Document()
    d.add_heading("Test Policy", 0)
    meta = d.add_table(rows=2, cols=2)
    for r, (k, v) in enumerate([("Document ID", "TEST-POL"), ("Version", "1.0")]):
        meta.cell(r, 0).text, meta.cell(r, 1).text = k, v
    d.add_heading("1. Leave rules", 1)
    d.add_paragraph("Staff get 7 test days each year. This line continues the same paragraph.")
    d.add_paragraph("A second paragraph starts here.")
    d.add_heading("2. Rates", 1)
    t = d.add_table(rows=3, cols=2)
    for r, row in enumerate([("Band", "Days"), ("A", "1.5"), ("B", "2.0")]):
        for c, value in enumerate(row):
            t.cell(r, c).text = value
    d.add_paragraph("bullet item", style="List Bullet")
    d.sections[0].footer.paragraphs[0].text = "Test Co | Page"
    path = tmp_path / "p.docx"
    d.save(path)

    blocks = load_document(str(path)).blocks
    assert _summary(blocks) == EXPECTED + [("text", 0, "- bullet item")]
    assert all(b.page is None for b in blocks)


HTML = """<!DOCTYPE html><html><head><title>ignored</title><style>h1{color:red}</style>
<script>var x = "not content";</script></head><body>
<nav>Intranet &rsaquo; Policies</nav>
<h1>Test Policy</h1>
<table class='meta'><tr><td>Document ID</td><td>TEST-POL</td></tr></table>
<h2>1. Leave rules</h2>
<p>Staff get 7 test days each year.
   This line continues the same paragraph.</p>
<div><p>A second paragraph starts here.</p></div>
<h2>2. Rates</h2>
<table><thead><tr><th>Band</th><th>Days</th></tr></thead>
<tbody><tr><td>A</td><td>1.5</td></tr><tr><td>B</td><td>2.0</td></tr></tbody></table>
<ul><li>one</li><li>two<ul><li>nested</li></ul></li></ul>
<footer>Test Co</footer></body></html>"""


def test_html(tmp_path):
    path = tmp_path / "p.html"
    path.write_text(HTML, encoding="utf-8")
    assert _summary(load_document(str(path)).blocks) == EXPECTED + [
        ("text", 0, "- one"),
        ("text", 0, "- two"),
        ("text", 0, "- nested"),
    ]


MARKDOWN = """# Test Policy

| Field | Value |
| --- | --- |
| Document ID | TEST-POL |

## 1. Leave rules

Staff get **7 test days** each year.
This line continues the same paragraph.

A second paragraph starts here.

## 2. Rates

| Band | Days |
| --- | --- |
| A | 1.5 |
| B | 2.0 |

- one
- two
"""


def test_markdown(tmp_path):
    path = tmp_path / "p.md"
    path.write_text(MARKDOWN, encoding="utf-8")
    assert _summary(load_document(str(path)).blocks) == EXPECTED + [("text", 0, "- one"), ("text", 0, "- two")]


def test_unsupported_extension(tmp_path):
    path = tmp_path / "p.txt"
    path.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError):
        load_document(str(path))
